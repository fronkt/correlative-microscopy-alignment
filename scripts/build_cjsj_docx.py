"""Build the CJSJ submission .docx from paper/cjsj/paper.md inside the journal's own Word template.

The template (paper/cjsj/templates/CJSJ-Original-Research-Template-1-faxy.docx) supplies page setup
(two columns, 0.75 in margins) and styles; this script empties its body and refills it.

paper.md conventions (deliberately tiny):
  line 1  "% Title"            line 2  "% Author line"      line 3  "% Affiliation line"
  "Abstract—..." paragraph     rendered with a bold-italic "Abstract—" lead-in
  "# Heading"                  section heading (template's Heading 1)
  "## Heading"                 subsection (italic run-in)
  "[FIG n]"                    placeholder line (figures travel in the .pptx; CJSJ counts them separately)
  "1. ..." after "# References" numbered references
  *italic* and **bold** inline
"""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "paper/cjsj/templates/CJSJ-Original-Research-Template-1-faxy.docx"
SRC = ROOT / "paper/cjsj/paper.md"
OUT = ROOT / "paper/cjsj/Cai_CJSJ_2026.docx"


def add_runs(p, text: str, size: float = 10, base_italic: bool = False) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        bold = part.startswith("**")
        ital = base_italic or (part.startswith("*") and not bold)
        r = p.add_run(part.strip("*") if (bold or part.startswith("*")) else part)
        r.bold, r.italic = bold or None, ital or None
        r.font.name, r.font.size = "Times New Roman", Pt(size)


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, spacing=0.95, first_indent=None):
    pf = p.paragraph_format
    pf.alignment, pf.space_after, pf.space_before, pf.line_spacing = align, Pt(after), Pt(0), spacing
    if first_indent is not None:
        pf.first_line_indent = Pt(first_indent)


def main() -> None:
    d = docx.Document(str(TEMPLATE))
    body = d.element.body
    sectpr = copy.deepcopy(body.sectPr)
    for child in list(body):
        body.remove(child)
    body.append(sectpr)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    title, author, affil = (l.lstrip("% ").strip() for l in lines[:3])
    p = d.add_paragraph(style="Title")
    add_runs(p, title, size=16)
    fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=4, spacing=1.0)
    for txt, it in ((author, False), (affil, True)):
        p = d.add_paragraph()
        add_runs(p, txt, size=10, base_italic=it)
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=2, spacing=1.0)

    in_refs, para = False, []

    def flush():
        nonlocal para
        if not para:
            return
        text = " ".join(para).strip()
        para = []
        if text.startswith("Abstract—"):
            p = d.add_paragraph()
            r = p.add_run("Abstract—")
            r.bold, r.italic = True, True
            r.font.name, r.font.size = "Times New Roman", Pt(9)
            add_runs(p, text[len("Abstract—"):], size=9)
            fmt(p, after=6)
        elif text.startswith("[FIG"):
            p = d.add_paragraph()
            add_runs(p, text, size=8, base_italic=True)
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=4)
        elif in_refs:
            m = re.match(r"(\d+)\.\s+(.*)", text)
            p = d.add_paragraph()
            add_runs(p, f"[{m.group(1)}] {m.group(2)}" if m else text, size=8)
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, after=1, spacing=1.0)
        else:
            p = d.add_paragraph()
            add_runs(p, text)
            fmt(p, first_indent=10)

    for line in lines[3:]:
        s = line.strip()
        if s.startswith("# "):
            flush()
            h = s[2:].strip()
            in_refs = h.lower() == "references"
            p = d.add_paragraph(style="Heading 1")
            add_runs(p, h.upper() if not in_refs else "References", size=10)
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=2, spacing=1.0)
        elif s.startswith("## "):
            flush()
            para = [f"*{s[3:].strip()}.*"]
        elif not s:
            flush()
        elif in_refs and re.match(r"\d+\.\s", s):
            flush()
            para = [s]
        else:
            para.append(s)
    flush()
    d.save(str(OUT))
    words = sum(len(re.sub(r"\*", "", l).split()) for l in lines[3:]
                if l.strip() and not l.startswith("#"))
    print(f"wrote {OUT} ({words} words incl. references)", file=sys.stderr)


if __name__ == "__main__":
    main()
